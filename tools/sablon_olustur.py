"""İş Takip Excel şablonunu (sablon/IsTakip.xlsx) oluşturur.

Kullanım:  pip install openpyxl && python tools/sablon_olustur.py
"""
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).resolve().parent.parent / "sablon" / "IsTakip.xlsx"

UNITS = [
    ("B01", "Satın Alma", "", ""),
    ("B02", "İnsan Kaynakları", "", ""),
    ("B03", "Bilgi İşlem", "", ""),
    ("B04", "Muhasebe", "", ""),
    ("B05", "İdari İşler", "", ""),
]
STATUSES = ["Atandı", "Devam Ediyor", "Tamamlandı", "İptal"]
PRIORITIES = ["Yüksek", "Orta", "Düşük"]
TASK_HEADERS = [
    "GörevNo", "AtamaTarihi", "Birim", "GörevTanımı", "Öncelik",
    "TerminTarihi", "Durum", "TamamlandıMı", "TamamlanmaTarihi", "Not",
]
MAX_ROWS = 2000
DATE_FMT = "DD.MM.YYYY"


def build():
    wb = Workbook()

    # --- Görevler -----------------------------------------------------------
    ws = wb.active
    ws.title = "Görevler"
    ws.append(TASK_HEADERS)
    today = date.today()
    ws.append(["G-0001", today, "Satın Alma", "Örnek görev — silebilirsiniz",
               "Orta", today + timedelta(days=7), "Atandı", False, None, ""])
    ws.add_table(Table(
        displayName="Gorevler", ref=f"A1:J{ws.max_row}",
        tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True),
    ))
    for col, width in zip("ABCDEFGHIJ", [10, 12, 18, 50, 10, 12, 14, 13, 16, 30]):
        ws.column_dimensions[col].width = width
    for col in "BFI":
        for row in range(2, MAX_ROWS + 1):
            ws[f"{col}{row}"].number_format = DATE_FMT
    ws.freeze_panes = "A2"

    def list_validation(formula, rng):
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.add(rng)
        ws.add_data_validation(dv)

    list_validation(f"Birimler!$B$2:$B${MAX_ROWS}", f"C2:C{MAX_ROWS}")
    list_validation('"' + ",".join(PRIORITIES) + '"', f"E2:E{MAX_ROWS}")
    list_validation('"' + ",".join(STATUSES) + '"', f"G2:G{MAX_ROWS}")

    rng = f"A2:J{MAX_ROWS}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=['$H2=TRUE'],
        font=Font(color="808080"), fill=PatternFill("solid", bgColor="E6F4EA")))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=['AND($H2<>TRUE,$F2<>"",$F2<TODAY(),$G2<>"İptal")'],
        font=Font(color="C62828"), fill=PatternFill("solid", bgColor="FDECEC")))

    # --- Birimler -----------------------------------------------------------
    wu = wb.create_sheet("Birimler")
    wu.append(["BirimKodu", "BirimAdı", "Sorumlu", "E-posta"])
    for u in UNITS:
        wu.append(list(u))
    wu.add_table(Table(
        displayName="Birimler", ref=f"A1:D{wu.max_row}",
        tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True),
    ))
    for col, width in zip("ABCD", [11, 24, 22, 28]):
        wu.column_dimensions[col].width = width

    # --- Özet ---------------------------------------------------------------
    so = wb.create_sheet("Özet")
    so.append(["Birim", "Toplam", "Tamamlanan", "Açık", "Geciken", "Oran"])
    for c in so[1]:
        c.font = Font(bold=True)
    for i, u in enumerate(UNITS, start=2):
        so.append([
            f"=Birimler!B{i}",
            f'=COUNTIFS(Gorevler[Birim],A{i},Gorevler[Durum],"<>İptal")',
            f'=COUNTIFS(Gorevler[Birim],A{i},Gorevler[TamamlandıMı],TRUE)',
            f"=B{i}-C{i}",
            f'=COUNTIFS(Gorevler[Birim],A{i},Gorevler[TamamlandıMı],"<>TRUE",'
            f'Gorevler[Durum],"<>İptal",Gorevler[TerminTarihi],"<"&TODAY())',
            f'=IF(B{i}=0,"",C{i}/B{i})',
        ])
        so[f"F{i}"].number_format = "0%"
    for col, width in zip("ABCDEF", [24, 10, 12, 8, 10, 8]):
        so.column_dimensions[col].width = width

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Oluşturuldu: {OUT}")


if __name__ == "__main__":
    build()
